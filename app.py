import bisect
import base64
import hashlib
import hmac
import html
import io
import json
import os
import re
import secrets
import sqlite3
import zipfile

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

# --- 全局与主题配置 (必须放在最顶部) ---
st.set_page_config(page_title="典藏传世名方系统V2.0", page_icon="📜", layout="wide")

# 注入自定义 CSS：庄重、古朴的中医古籍风格
st.markdown("""
    <style>
    .stApp { background-color: #FDFBF7; font-family: 'Noto Serif SC', 'SimSun', 'Songti SC', serif; color: #2C2825; }
    h1, h2, h3 { color: #792D11 !important; font-weight: bold; }
    [data-testid="stSidebar"] { background-color: #F2E8D5; border-right: 2px solid #D9C5A0; }
    .stButton>button { background-color: #792D11; color: #FFFFFF; border-radius: 4px; border: 1px solid #4A1A08; transition: all 0.3s ease; font-weight: bold; }
    .stButton>button:hover { background-color: #4A1A08; color: #FDFBF7; border: 1px solid #792D11; }
    [data-testid="stVerticalBlockBorderWrapper"] { background-color: #FFFFFF; border: 1px solid #E8DCC4 !important; box-shadow: 0 4px 6px rgba(0,0,0,0.02); border-radius: 8px; }
    mark.yellow-highlight { background-color: #FFF2CC; color: #B83B00; font-weight: bold; padding: 0 4px; border-radius: 2px; border-bottom: 2px solid #F0B27A; }
    .book-reader { max-width: 920px; margin: 18px auto 30px; background: #FFFDF8; border: 1px solid #E8DCC4; border-radius: 12px; padding: 38px 48px; box-shadow: 0 8px 24px rgba(90,60,30,0.08); font-family: 'Noto Serif SC', 'SimSun', 'Songti SC', serif; line-height: 2.05; color: #2C2825; text-align: justify; }
    .book-reader p { text-indent: 2em; margin: 0 0 0.9em; }
    @keyframes pageFade { from { opacity: 0.25; transform: translateX(8px); } to { opacity: 1; transform: none; } }
    .page-fade { animation: pageFade 0.18s ease-out; }
    @media (max-width: 700px) { .book-reader { padding: 22px 20px; } }
    .disclaimer { color: #B83B00; font-size: 0.85rem; padding: 10px; border-top: 1px dashed #D9C5A0; margin-top: 30px; }
    </style>
""", unsafe_allow_html=True)

DB_FILE = "tcm_database.db"
ZIP_FILE = "tcm_database.zip"
DB_SCHEMA_VERSION = 5  # 数据库结构变更时递增，使 Streamlit 资源缓存失效并重新执行迁移
VERIF_PAGE_SIZE = 20  # 各种验方检索每页条数
FORMULA_PAGE_SIZE = 20  # 传世名方基础检索每页条数
BOOK_PAGE_CHARS = 1200  # 图书文本阅读每页字符数
BOOK_BACKUP_FILE = os.environ.get(
    "TCM_BOOK_BACKUP_FILE",
    os.path.join(os.path.dirname(os.path.abspath(DB_FILE)), "book_library_backup.db"))

BOOK_CATEGORIES = [
    "基础理论类",
    "临床实践类",
    "方药本草类",
    "医案医话类",
    "历史文献类",
    "食疗养生类",
]
BOOK_FILE_TYPES = ["pdf", "doc", "docx", "wps", "txt", "md", "rtf"]

FORMULA_TABLE = "formulas"
RECIPE_TABLE = "recipes"
VERIF_TABLE = "各种验方表"

VERIF_COLUMNS = ["方名", "方剂及用法", "主治疾病", "症状", "案例", "出处"]
FORMULA_COLUMNS = [
    "id", "大医名", "方剂名", "方剂详情", "临床病名", "病例详情",
    "按语", "现代研究", "临证提要", "方别", "其他",
]
RECIPE_COLUMNS = [
    "id", "菜名", "材料和制作", "功效", "注意事项", "适合季节",
    "适合年纪", "适合性别", "餐别", "版别", "出处", "编者及其身份", "品类",
]
TABLE_COLUMNS = {
    FORMULA_TABLE: FORMULA_COLUMNS,
    RECIPE_TABLE: RECIPE_COLUMNS,
}

DEFAULT_AI_CONFIG = {
    "api_key": "",
    "base_url": "https://api.deepseek.com/v1",
    "model_name": "deepseek-chat",
}

# AI 用户设置的浏览器 localStorage 存储组件
AI_STORAGE_COMPONENT_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "ai_storage_component")

# 病名编号：阿拉伯数字、全角数字、中文数字
_NUM_CHARS = "0-9０-９一二三四五六七八九十百零〇两"
# （一） / (1) / （1） 等括号编号
_PAREN_NUM_RE = re.compile(rf"^[（(]\s*[{_NUM_CHARS}\s]+?\s*[）)]\s*")
# 1. / 1、 / 1) / 一、 / 三． 等编号
_DOT_NUM_RE = re.compile(rf"^[{_NUM_CHARS}]+\s*[.．、,，)）:：]\s*")
# 数据中偶见的缺左括号写法，如 “四)神经精神疾病”
_UNBALANCED_PAREN_RE = re.compile(rf"^[（(]?\s*[{_NUM_CHARS}\s]+?\s*[）)]\s*")
# 去掉编号后可能残留的前导标点，如 “29..前列腺增生症”
_LEADING_PUNCT_RE = re.compile(r"^[.．、,，;；:：\-—_·\s]+")


# --- 1. 数据库解压与初始化 ---
def get_table_columns(conn, table_name):
    """读取数据表列名（中文表名同样适用）。"""
    return [row[1] for row in conn.execute(f'PRAGMA table_info("{table_name}")')]


def migrate_recipes_table(conn):
    """把旧版食谱表迁移为：功效、注意事项 两列，且注意事项位于功效之后、适合季节之前。"""
    cols = get_table_columns(conn, RECIPE_TABLE)
    if not cols:
        return
    if cols == RECIPE_COLUMNS:
        return

    def pick(*names):
        for name in names:
            if name in cols:
                return f'"{name}"'
        return "NULL"

    conn.execute("DROP TABLE IF EXISTS recipes__legacy_migration")
    conn.execute("ALTER TABLE recipes RENAME TO recipes__legacy_migration")
    conn.execute("""
        CREATE TABLE recipes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            菜名 TEXT,
            材料和制作 TEXT,
            功效 TEXT,
            注意事项 TEXT,
            适合季节 TEXT,
            适合年纪 TEXT,
            适合性别 TEXT,
            餐别 TEXT,
            版别 TEXT,
            出处 TEXT,
            编者及其身份 TEXT,
            品类 TEXT
        )
    """)
    select_expr = ", ".join([
        pick("id"),
        pick("菜名"),
        pick("材料和制作"),
        pick("功效", "功效和注意事项"),
        pick("注意事项", "季节功能"),
        pick("适合季节"),
        pick("适合年纪"),
        pick("适合性别"),
        pick("餐别"),
        pick("版别"),
        pick("出处"),
        pick("编者及其身份"),
        pick("品类"),
    ])
    conn.execute(f"""
        INSERT INTO recipes ({", ".join(f'"{c}"' for c in RECIPE_COLUMNS)})
        SELECT {select_expr} FROM recipes__legacy_migration
    """)
    conn.execute("DROP TABLE recipes__legacy_migration")


def migrate_verif_table(conn):
    """把旧版各种验方表迁移为《各种验方数据表（式样）》的六列结构。"""
    cols = get_table_columns(conn, VERIF_TABLE)
    desired = ["id"] + VERIF_COLUMNS
    if not cols or cols == desired:
        return

    def pick(*names):
        for name in names:
            if name in cols:
                return f'"{name}"'
        return "NULL"

    case_expr = pick("案例")
    if "备注" in cols:
        base_expr = 'COALESCE("案例", \'\')' if "案例" in cols else "''"
        case_expr = (
            f"CASE WHEN \"备注\" IS NOT NULL AND TRIM(\"备注\") <> '' "
            f"THEN {base_expr} || CASE WHEN {base_expr} = '' THEN '' ELSE char(10) END "
            f"|| '【备注】' || \"备注\" ELSE {base_expr} END"
        )

    temp_name = f"{VERIF_TABLE}__legacy_migration"
    conn.execute(f'DROP TABLE IF EXISTS "{temp_name}"')
    conn.execute(f'ALTER TABLE "{VERIF_TABLE}" RENAME TO "{temp_name}"')
    conn.execute(f'''CREATE TABLE "{VERIF_TABLE}"
                 (id INTEGER PRIMARY KEY AUTOINCREMENT,
                  方名 TEXT,
                  方剂及用法 TEXT,
                  主治疾病 TEXT,
                  症状 TEXT,
                  案例 TEXT,
                  出处 TEXT)''')
    select_expr = ", ".join([
        pick("id"),
        pick("方名"),
        pick("方剂及用法", "方剂"),
        pick("主治疾病", "主治功能"),
        pick("症状"),
        case_expr,
        pick("出处", "来源"),
    ])
    conn.execute(f"""
        INSERT INTO "{VERIF_TABLE}" ({", ".join(f'"{c}"' for c in desired)})
        SELECT {select_expr} FROM "{temp_name}"
    """)
    conn.execute(f'DROP TABLE "{temp_name}"')


ADMIN_PASSWORD_SALT_KEY = "admin_password_salt"
ADMIN_PASSWORD_HASH_KEY = "admin_password_hash"
ADMIN_PASSWORD_ITERATIONS = 120000
DEFAULT_ADMIN_PASSWORD = "123"


def hash_admin_password(password, salt_hex):
    return hashlib.pbkdf2_hmac(
        "sha256",
        str(password).encode("utf-8"),
        bytes.fromhex(salt_hex),
        ADMIN_PASSWORD_ITERATIONS,
    ).hex()


def ensure_admin_password(conn):
    """确保密码配置表存在；首次使用时以默认密码 123 初始化。"""
    conn.execute(
        "CREATE TABLE IF NOT EXISTS system_config (key TEXT PRIMARY KEY, value TEXT)")
    rows = dict(conn.execute("SELECT key, value FROM system_config").fetchall())
    salt = rows.get(ADMIN_PASSWORD_SALT_KEY)
    pwd_hash = rows.get(ADMIN_PASSWORD_HASH_KEY)
    if not salt or not pwd_hash:
        salt = secrets.token_hex(16)
        pwd_hash = hash_admin_password(DEFAULT_ADMIN_PASSWORD, salt)
        conn.execute(
            "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
            (ADMIN_PASSWORD_SALT_KEY, salt))
        conn.execute(
            "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
            (ADMIN_PASSWORD_HASH_KEY, pwd_hash))
    return salt, pwd_hash


def rebuild_imported_table(conn, table_name, columns):
    """Excel 导入后重建标准表结构，恢复 id 主键与自增，避免后续增删改失效。"""
    existing = get_table_columns(conn, table_name)
    if not existing:
        return
    temp_name = f"{table_name}__imported"
    conn.execute(f'DROP TABLE IF EXISTS "{temp_name}"')
    conn.execute(f'ALTER TABLE "{table_name}" RENAME TO "{temp_name}"')
    col_defs = ", ".join(f'"{c}" TEXT' for c in columns if c != "id")
    conn.execute(f'CREATE TABLE "{table_name}" (id INTEGER PRIMARY KEY AUTOINCREMENT, {col_defs})')
    select_expr = ", ".join(f'"{c}"' if c in existing else "NULL" for c in columns)
    col_list = ", ".join(f'"{c}"' for c in columns)
    conn.execute(
        f'INSERT INTO "{table_name}" ({col_list}) SELECT {select_expr} FROM "{temp_name}"')
    conn.execute(f'DROP TABLE "{temp_name}"')


def ensure_schema(conn):
    """建表与旧库迁移，可重复执行。"""
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS query_logs
                 (id INTEGER PRIMARY KEY AUTOINCREMENT,
                  ip_address TEXT, search_type TEXT, keyword TEXT,
                  query_time DATETIME DEFAULT CURRENT_TIMESTAMP)''')

    c.execute('''CREATE TABLE IF NOT EXISTS ai_config
                 (id INTEGER PRIMARY KEY CHECK (id = 1),
                  api_key TEXT, base_url TEXT, model_name TEXT)''')

    c.execute("SELECT COUNT(*) FROM ai_config")
    if c.fetchone()[0] == 0:
        c.execute(
            "INSERT INTO ai_config (id, api_key, base_url, model_name) VALUES (1, '', 'https://api.deepseek.com/v1', 'deepseek-chat')")

    c.execute(f'''CREATE TABLE IF NOT EXISTS "{VERIF_TABLE}"
                 (id INTEGER PRIMARY KEY AUTOINCREMENT,
                  方名 TEXT,
                  方剂及用法 TEXT,
                  主治疾病 TEXT,
                  症状 TEXT,
                  案例 TEXT,
                  出处 TEXT)''')

    c.execute('''CREATE TABLE IF NOT EXISTS recipes
                 (id INTEGER PRIMARY KEY AUTOINCREMENT,
                  菜名 TEXT,
                  材料和制作 TEXT,
                  功效 TEXT,
                  注意事项 TEXT,
                  适合季节 TEXT,
                  适合年纪 TEXT,
                  适合性别 TEXT,
                  餐别 TEXT,
                  版别 TEXT,
                  出处 TEXT,
                  编者及其身份 TEXT,
                  品类 TEXT)''')

    c.execute('''CREATE TABLE IF NOT EXISTS books
                 (id INTEGER PRIMARY KEY AUTOINCREMENT,
                  category TEXT NOT NULL,
                  title TEXT NOT NULL,
                  format TEXT,
                  filename TEXT,
                  text_content TEXT,
                  char_count INTEGER DEFAULT 0,
                  page_count INTEGER DEFAULT 0,
                  file_size INTEGER DEFAULT 0,
                  uploaded_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                  UNIQUE(category, title))''')

    c.execute('''CREATE TABLE IF NOT EXISTS book_files
                 (book_id INTEGER PRIMARY KEY,
                  content BLOB)''')

    c.execute('''CREATE TABLE IF NOT EXISTS reading_progress
                 (book_id INTEGER PRIMARY KEY,
                  position INTEGER DEFAULT 0,
                  updated_at DATETIME DEFAULT CURRENT_TIMESTAMP)''')

    c.execute('''CREATE INDEX IF NOT EXISTS idx_books_category
                 ON books(category, title)''')

    migrate_recipes_table(conn)
    migrate_verif_table(conn)
    ensure_admin_password(conn)
    conn.commit()


@st.cache_resource
def init_system(schema_version=DB_SCHEMA_VERSION):
    # schema_version 作为缓存键：数据库结构升级后即使热重载也会重新执行迁移
    _ = schema_version
    if not os.path.exists(DB_FILE):
        if os.path.exists(ZIP_FILE):
            try:
                with zipfile.ZipFile(ZIP_FILE, 'r') as zip_ref:
                    zip_ref.extractall(".")
            except Exception as e:
                return False, f"解压失败: {e}"
        else:
            return False, "未找到数据库文件或压缩包。"

    conn = sqlite3.connect(DB_FILE)
    try:
        ensure_schema(conn)
    except Exception as e:
        conn.close()
        return False, f"数据库结构初始化失败: {e}"
    conn.close()
    return True, "Success"


status, msg = init_system(DB_SCHEMA_VERSION)
if not status:
    st.error(f"⚠️️ 系统初始化失败: {msg}")
    st.stop()


def get_connection():
    return sqlite3.connect(DB_FILE, check_same_thread=False)


def verify_admin_password(password):
    """校验后台密码；数据库结构异常时按验证失败处理。"""
    try:
        conn = get_connection()
        try:
            salt, pwd_hash = ensure_admin_password(conn)
            conn.commit()
        finally:
            conn.close()
        return hmac.compare_digest(hash_admin_password(password, salt), pwd_hash)
    except Exception:
        return False


def set_admin_password(new_password):
    """更新后台密码，返回 (是否成功, 错误信息)。"""
    try:
        conn = get_connection()
        try:
            ensure_admin_password(conn)
            salt = secrets.token_hex(16)
            pwd_hash = hash_admin_password(new_password, salt)
            conn.execute(
                "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
                (ADMIN_PASSWORD_SALT_KEY, salt))
            conn.execute(
                "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
                (ADMIN_PASSWORD_HASH_KEY, pwd_hash))
            conn.commit()
        finally:
            conn.close()
        return True, ""
    except Exception as e:
        return False, str(e)


DOWNLOAD_PASSWORD_SALT_KEY = "download_password_salt"
DOWNLOAD_PASSWORD_HASH_KEY = "download_password_hash"
DEFAULT_DOWNLOAD_PASSWORD = "123"


def ensure_download_password(conn):
    """确保下载密码已初始化，默认 123，可在后台修改。"""
    conn.execute(
        "CREATE TABLE IF NOT EXISTS system_config (key TEXT PRIMARY KEY, value TEXT)")
    rows = dict(conn.execute("SELECT key, value FROM system_config").fetchall())
    salt = rows.get(DOWNLOAD_PASSWORD_SALT_KEY)
    pwd_hash = rows.get(DOWNLOAD_PASSWORD_HASH_KEY)
    if not salt or not pwd_hash:
        salt = secrets.token_hex(16)
        pwd_hash = hash_admin_password(DEFAULT_DOWNLOAD_PASSWORD, salt)
        conn.execute(
            "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
            (DOWNLOAD_PASSWORD_SALT_KEY, salt))
        conn.execute(
            "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
            (DOWNLOAD_PASSWORD_HASH_KEY, pwd_hash))
    return salt, pwd_hash


def verify_download_password(password):
    try:
        conn = get_connection()
        try:
            salt, pwd_hash = ensure_download_password(conn)
            conn.commit()
        finally:
            conn.close()
        return hmac.compare_digest(hash_admin_password(password, salt), pwd_hash)
    except Exception:
        return False


def set_download_password(new_password):
    try:
        conn = get_connection()
        try:
            ensure_download_password(conn)
            salt = secrets.token_hex(16)
            pwd_hash = hash_admin_password(new_password, salt)
            conn.execute(
                "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
                (DOWNLOAD_PASSWORD_SALT_KEY, salt))
            conn.execute(
                "INSERT OR REPLACE INTO system_config (key, value) VALUES (?, ?)",
                (DOWNLOAD_PASSWORD_HASH_KEY, pwd_hash))
            conn.commit()
        finally:
            conn.close()
        return True, ""
    except Exception as e:
        return False, str(e)


@st.cache_resource
def _get_ai_storage_component():
    """浏览器 localStorage 存储组件；组件目录缺失时返回 None，页面仍可正常使用。"""
    if not os.path.isdir(AI_STORAGE_COMPONENT_DIR):
        return None
    try:
        return components.declare_component("tcm_ai_storage", path=AI_STORAGE_COMPONENT_DIR)
    except Exception:
        return None


def _save_ai_settings_callback():
    """保存按钮回调：把当前输入交给浏览器存储组件写入 localStorage。"""
    config = {
        "api_key": str(st.session_state.get("ai_user_api_key", "")).strip(),
        "base_url": str(st.session_state.get("ai_user_base_url", "")).strip(),
        "model_name": str(st.session_state.get("ai_user_model_name", "")).strip(),
    }
    st.session_state.ai_user_config = config
    st.session_state["ai_storage_action"] = "save"
    st.session_state["ai_storage_payload"] = config
    st.session_state["ai_storage_action_id"] = int(
        st.session_state.get("ai_storage_action_id", 0)) + 1


def _clear_ai_settings_callback():
    """清空按钮回调：让浏览器组件删除 localStorage，并重置输入框。"""
    st.session_state.ai_user_config = dict(DEFAULT_AI_CONFIG)
    st.session_state["ai_user_api_key"] = ""
    st.session_state["ai_user_base_url"] = DEFAULT_AI_CONFIG["base_url"]
    st.session_state["ai_user_model_name"] = DEFAULT_AI_CONFIG["model_name"]
    st.session_state["ai_storage_action"] = "clear"
    st.session_state["ai_storage_payload"] = None
    st.session_state["ai_storage_action_id"] = int(
        st.session_state.get("ai_storage_action_id", 0)) + 1


def _dismiss_ai_settings_notice():
    """用户重新编辑设置时，移除上一次的保存/清空提示。"""
    st.session_state.pop("ai_settings_notice", None)


def log_query(search_type, keyword):
    try:
        conn = get_connection()
        c = conn.cursor()
        c.execute("INSERT INTO query_logs (search_type, keyword) VALUES (?, ?)", (search_type, keyword))
        conn.commit()
        conn.close()
    except Exception:
        pass


def highlight_text(text, kw_list):
    if not isinstance(text, str) or not text:
        return ""
    for kw in kw_list:
        if kw and len(kw) > 1:
            text = text.replace(kw, f"<mark class='yellow-highlight'>{kw}</mark>")
    return text


def safe_text(value, default=""):
    if value is None:
        return default
    try:
        if pd.isna(value):
            return default
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def strip_leading_number(text):
    """去掉字符串开头的层级编号，如 (一)、1.、一、、四)。"""
    s = str(text or "").strip()
    for _ in range(4):
        new = _PAREN_NUM_RE.sub("", s, count=1)
        if new == s:
            new = _DOT_NUM_RE.sub("", s, count=1)
        if new == s:
            new = _UNBALANCED_PAREN_RE.sub("", s, count=1)
        if new == s:
            break
        s = new.strip()
    s = _LEADING_PUNCT_RE.sub("", s)
    return s.strip()


def normalize_disease_name(raw):
    """取多层病名的最后一层，去掉前面的编号，并清理空白。"""
    if raw is None:
        return ""
    try:
        if pd.isna(raw):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(raw).strip().replace("\u3000", " ")
    if not s:
        return ""
    parts = [p.strip() for p in re.split(r"[/／]", s) if p.strip()]
    if not parts:
        return ""
    s = strip_leading_number(parts[-1])
    s = re.sub(r"\s+", "", s)
    return s.strip()


def render_formula_card(row):
    with st.container(border=True):
        st.markdown(f"### 🛡️ 【{safe_text(row.get('方剂名'), '未知')}】")
        st.markdown(
            f"**👨‍⚕️ 医家名:** {safe_text(row.get('大医名'), '未知')} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"**🏷️ 方别:** {safe_text(row.get('方别'), '未知')} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"**🩺 临床病名:** {safe_text(row.get('临床病名'), '未知')}")
        st.markdown(f"**📜 方剂详情:**<br>{safe_text(row.get('方剂详情'))}", unsafe_allow_html=True)
        st.markdown(f"**📝 临证提要:**<br>{safe_text(row.get('临证提要'), '无')}", unsafe_allow_html=True)
        st.markdown(f"**📖 病例详情:**<br>{safe_text(row.get('病例详情'), '无')}", unsafe_allow_html=True)
        if safe_text(row.get('按语')):
            st.markdown(f"**💡 按语:**<br>{safe_text(row.get('按语'))}", unsafe_allow_html=True)
        if safe_text(row.get('现代研究')):
            st.markdown(f"**🔬 现代研究:**<br>{safe_text(row.get('现代研究'))}", unsafe_allow_html=True)


def render_verif_card(row):
    with st.container(border=True):
        st.markdown(f"### 🧾 【{safe_text(row.get('方名'), '未命名')}】")
        st.markdown(f"**📜 方剂及用法:**<br>{safe_text(row.get('方剂及用法'), '无')}", unsafe_allow_html=True)
        st.markdown(f"**🩺 主治疾病:**<br>{safe_text(row.get('主治疾病'), '无')}", unsafe_allow_html=True)
        st.markdown(f"**🤒 症状:**<br>{safe_text(row.get('症状'), '无')}", unsafe_allow_html=True)
        st.markdown(f"**📖 案例:**<br>{safe_text(row.get('案例'), '无')}", unsafe_allow_html=True)
        st.markdown(f"**📚 出处:** {safe_text(row.get('出处'), '未注明')}")


def split_verif_indications(raw):
    """把“主治疾病”拆成可选的病名项，兼容多行、分号、顿号。"""
    if raw is None:
        return []
    try:
        if pd.isna(raw):
            return []
    except (TypeError, ValueError):
        pass
    parts = re.split(r"[\r\n；;、/／]+", str(raw))
    return [p.strip() for p in parts if p.strip()]


def format_file_size(size):
    try:
        size = int(size or 0)
    except (TypeError, ValueError):
        return "0 B"
    if size < 1024:
        return f"{size} B"
    size /= 1024
    if size < 1024:
        return f"{size:.1f} KB"
    size /= 1024
    if size < 1024:
        return f"{size:.1f} MB"
    return f"{size / 1024:.1f} GB"


def decode_book_bytes(data):
    """尽量以中文常见编码解码 TXT/类文本文件。"""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="ignore")


def extract_book_content(data, filename):
    """提取图书文字，返回 (文本, PDF页数, 错误信息)。"""
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if ext == "pdf":
        try:
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
            texts = []
            for page in reader.pages:
                try:
                    texts.append(page.extract_text() or "")
                except Exception:
                    texts.append("")
            text = "\n\n".join(texts).replace("\r\n", "\n").replace("\r", "\n").strip()
            return text, len(reader.pages), ""
        except Exception as e:
            return "", 0, f"PDF 解析失败：{e}"

    if ext in ("docx", "doc", "wps"):
        if data[:2] != b"PK":
            return "", 0, "该文件为旧版 DOC/WPS 二进制格式，暂不支持提取文字；原文件已保存，可下载阅读。"
        try:
            from docx import Document
            document = Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
            text = "\n".join(parts).strip()
            return text, 0, ""
        except Exception as e:
            return "", 0, f"Word/WPS 解析失败：{e}"

    text = decode_book_bytes(data).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text, 0, ""


def upsert_book(conn, category, title, fmt, filename, data, text, page_count):
    """按分类+书名新增或覆盖图书，保留已有阅读进度。"""
    existing = conn.execute(
        "SELECT id FROM books WHERE category=? AND title=?", (category, title)).fetchone()
    size = len(data)
    if existing:
        book_id = int(existing[0])
        conn.execute(
            "UPDATE books SET format=?, filename=?, text_content=?, char_count=?, "
            "page_count=?, file_size=?, uploaded_at=CURRENT_TIMESTAMP WHERE id=?",
            (fmt, filename, text, len(text), int(page_count or 0), size, book_id))
        is_new = False
    else:
        cur = conn.execute(
            "INSERT INTO books (category, title, format, filename, text_content, "
            "char_count, page_count, file_size) VALUES (?,?,?,?,?,?,?,?)",
            (category, title, fmt, filename, text, len(text), int(page_count or 0), size))
        book_id = int(cur.lastrowid)
        is_new = True
    conn.execute(
        "INSERT OR REPLACE INTO book_files (book_id, content) VALUES (?, ?)",
        (book_id, sqlite3.Binary(data)))
    return book_id, is_new


@st.cache_data(show_spinner=False)
def load_books(category):
    conn = get_connection()
    try:
        return pd.read_sql(
            "SELECT id, title, format, filename, file_size, char_count, page_count "
            "FROM books WHERE category=? ORDER BY title",
            conn, params=(category,))
    finally:
        conn.close()


@st.cache_data(show_spinner=False)
def load_book_text(book_id):
    conn = get_connection()
    try:
        return conn.execute(
            "SELECT title, format, text_content, char_count, page_count, file_size "
            "FROM books WHERE id=?", (book_id,)).fetchone()
    finally:
        conn.close()


def load_book_file(book_id):
    conn = get_connection()
    try:
        row = conn.execute(
            "SELECT content FROM book_files WHERE book_id=?", (book_id,)).fetchone()
        return row[0] if row else None
    finally:
        conn.close()


def get_pdf_page_count(data):
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(data)
        try:
            return len(pdf)
        finally:
            pdf.close()
    except Exception:
        return 0


def get_pdf_document(book_id, data):
    """缓存 PDF 文档对象，翻页/缩放时不再反复解析几十 MB 的文件。"""
    key = f"pdf_doc_{book_id}"
    doc = st.session_state.get(key)
    if doc is None:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(data)
        st.session_state[key] = doc
    return doc


def render_pdf_page_cached(book_id, data, page_index):
    """按页缓存 JPEG，缩放只改 CSS 宽度，不重新渲染 PDF。"""
    cache_key = f"pdf_page_cache_{book_id}"
    cache = st.session_state.get(cache_key)
    if cache is None:
        cache = {}
        st.session_state[cache_key] = cache
    if page_index in cache:
        return cache[page_index]

    page = None
    bitmap = None
    image = None
    try:
        doc = get_pdf_document(book_id, data)
        if page_index < 0 or page_index >= len(doc):
            return None
        page = doc[page_index]
        bitmap = page.render(scale=1.7)
        image = bitmap.to_pil()
        buffer = io.BytesIO()
        image.save(buffer, format="JPEG", quality=82, optimize=True)
        result = buffer.getvalue()
        if len(cache) >= 24:
            cache.clear()
        cache[page_index] = result
        return result
    except Exception:
        return None
    finally:
        if image is not None:
            try:
                image.close()
            except Exception:
                pass
        if bitmap is not None:
            try:
                bitmap.close()
            except Exception:
                pass
        if page is not None:
            try:
                page.close()
            except Exception:
                pass


def get_reading_position(book_id):
    try:
        conn = get_connection()
        try:
            row = conn.execute(
                "SELECT position FROM reading_progress WHERE book_id=?", (book_id,)).fetchone()
            return int(row[0]) if row and row[0] is not None else 0
        finally:
            conn.close()
    except Exception:
        return 0


def save_reading_position(book_id, position):
    try:
        conn = get_connection()
        try:
            conn.execute(
                "INSERT OR REPLACE INTO reading_progress (book_id, position, updated_at) "
                "VALUES (?, ?, CURRENT_TIMESTAMP)",
                (book_id, max(0, int(position))))
            conn.commit()
        finally:
            conn.close()
        save_backup_progress(book_id, position)
    except Exception:
        pass


BOOK_BACKUP_SCHEMA = """
CREATE TABLE IF NOT EXISTS books (
  id INTEGER PRIMARY KEY, category TEXT NOT NULL, title TEXT NOT NULL,
  format TEXT, filename TEXT, text_content TEXT, char_count INTEGER DEFAULT 0,
  page_count INTEGER DEFAULT 0, file_size INTEGER DEFAULT 0,
  uploaded_at DATETIME DEFAULT CURRENT_TIMESTAMP, UNIQUE(category, title));
CREATE TABLE IF NOT EXISTS book_files (book_id INTEGER PRIMARY KEY, content BLOB);
CREATE TABLE IF NOT EXISTS reading_progress (
  book_id INTEGER PRIMARY KEY, position INTEGER DEFAULT 0,
  updated_at DATETIME DEFAULT CURRENT_TIMESTAMP);
"""


def backup_book_library():
    """把图书库单独备份，避免部署包更新主数据库时丢失图书。"""
    try:
        if not os.path.exists(DB_FILE):
            return
        backup_dir = os.path.dirname(os.path.abspath(BOOK_BACKUP_FILE))
        if backup_dir:
            os.makedirs(backup_dir, exist_ok=True)
        src = sqlite3.connect(DB_FILE)
        dst = sqlite3.connect(BOOK_BACKUP_FILE)
        try:
            dst.executescript(BOOK_BACKUP_SCHEMA)
            dst.execute("DELETE FROM reading_progress")
            dst.execute("DELETE FROM book_files")
            dst.execute("DELETE FROM books")
            for row in src.execute(
                    "SELECT id, category, title, format, filename, text_content, "
                    "char_count, page_count, file_size, uploaded_at FROM books"):
                dst.execute(
                    "INSERT INTO books VALUES (?,?,?,?,?,?,?,?,?,?)", row)
            for row in src.execute("SELECT book_id, content FROM book_files"):
                dst.execute("INSERT INTO book_files VALUES (?,?)", row)
            for row in src.execute(
                    "SELECT book_id, position, updated_at FROM reading_progress"):
                dst.execute("INSERT INTO reading_progress VALUES (?,?,?)", row)
            dst.commit()
        finally:
            dst.close()
            src.close()
    except Exception:
        pass


def save_backup_progress(book_id, position):
    if not os.path.exists(BOOK_BACKUP_FILE):
        return
    try:
        conn = sqlite3.connect(BOOK_BACKUP_FILE)
        try:
            conn.execute("CREATE TABLE IF NOT EXISTS reading_progress "
                         "(book_id INTEGER PRIMARY KEY, position INTEGER DEFAULT 0, "
                         "updated_at DATETIME DEFAULT CURRENT_TIMESTAMP)")
            conn.execute(
                "INSERT OR REPLACE INTO reading_progress "
                "(book_id, position, updated_at) VALUES (?, ?, CURRENT_TIMESTAMP)",
                (int(book_id), max(0, int(position))))
            conn.commit()
        finally:
            conn.close()
    except Exception:
        pass


def restore_book_library_from_backup(force=False):
    if not os.path.exists(BOOK_BACKUP_FILE) or not os.path.exists(DB_FILE):
        return False
    try:
        src = sqlite3.connect(BOOK_BACKUP_FILE)
        try:
            backup_count = src.execute("SELECT COUNT(*) FROM books").fetchone()[0]
        finally:
            src.close()
        if not backup_count:
            return False
        main = get_connection()
        try:
            if main.execute("SELECT COUNT(*) FROM books").fetchone()[0] and not force:
                return False
            src = sqlite3.connect(BOOK_BACKUP_FILE)
            try:
                main.execute("DELETE FROM reading_progress")
                main.execute("DELETE FROM book_files")
                main.execute("DELETE FROM books")
                for row in src.execute(
                        "SELECT id, category, title, format, filename, text_content, "
                        "char_count, page_count, file_size, uploaded_at FROM books"):
                    main.execute("INSERT INTO books VALUES (?,?,?,?,?,?,?,?,?,?)", row)
                for row in src.execute("SELECT book_id, content FROM book_files"):
                    main.execute("INSERT INTO book_files VALUES (?,?)", row)
                for row in src.execute(
                        "SELECT book_id, position, updated_at FROM reading_progress"):
                    main.execute("INSERT INTO reading_progress VALUES (?,?,?)", row)
                main.commit()
                return True
            finally:
                src.close()
        finally:
            main.close()
    except Exception:
        return False


def bootstrap_book_library():
    """主库为空时从备份恢复；主库有图书而备份为空时先备份。"""
    try:
        if not os.path.exists(DB_FILE):
            return
        main_count = 0
        conn = sqlite3.connect(DB_FILE)
        try:
            main_count = conn.execute("SELECT COUNT(*) FROM books").fetchone()[0]
        finally:
            conn.close()
        backup_count = 0
        if os.path.exists(BOOK_BACKUP_FILE):
            try:
                b = sqlite3.connect(BOOK_BACKUP_FILE)
                try:
                    backup_count = b.execute("SELECT COUNT(*) FROM books").fetchone()[0]
                finally:
                    b.close()
            except Exception:
                backup_count = 0
        if main_count and not backup_count:
            backup_book_library()
        elif not main_count and backup_count:
            restore_book_library_from_backup()
    except Exception:
        pass


def build_book_pages(text):
    pages = []
    n = len(text)
    start = 0
    while start < n:
        end = min(start + BOOK_PAGE_CHARS, n)
        if end < n:
            cut = text.rfind("\n", start + int(BOOK_PAGE_CHARS * 0.6), end)
            if cut > start:
                end = cut + 1
        pages.append((start, text[start:end]))
        start = end
    return pages


def book_page_html(text, font_size):
    escaped = html.escape(text)
    blocks = re.split(r"\n\s*\n", escaped)
    parts = []
    for block in blocks:
        block = block.strip("\n")
        if not block.strip():
            continue
        parts.append(f"<p>{block.replace(chr(10), '<br>')}</p>")
    return (
        f"<div class='book-reader page-fade' style='font-size:{int(font_size)}px'>"
        + "".join(parts)
        + "</div>"
    )


def render_book_nav(book_id, page_index, pages, suffix, force_key=None):
    prev_col, next_col, info_col = st.columns([1, 1, 3])
    with prev_col:
        if st.button("⬅️ 上一页", disabled=page_index <= 0,
                     key=f"book_prev_{suffix}_{book_id}"):
            save_reading_position(book_id, pages[page_index - 1][0])
            if force_key:
                st.session_state[force_key] = True
            st.rerun()
    with next_col:
        if st.button("下一页 ➡️", disabled=page_index >= len(pages) - 1,
                     key=f"book_next_{suffix}_{book_id}"):
            save_reading_position(book_id, pages[page_index + 1][0])
            if force_key:
                st.session_state[force_key] = True
            st.rerun()
    with info_col:
        st.markdown(
            f"<div style='text-align:right;padding-top:6px'>第 {page_index + 1} / {len(pages)} 页</div>",
            unsafe_allow_html=True)


def render_protected_download(book_id, filename, content, mime):
    """下载原文件前要求输入下载密码。"""
    unlock_key = f"download_unlocked_{book_id}"
    if st.session_state.get(unlock_key):
        st.download_button(
            "⬇️ 下载原文件",
            data=content,
            file_name=filename,
            mime=mime,
            key=f"book_download_{book_id}")
        if st.button("🔒 锁定下载", key=f"download_lock_{book_id}"):
            st.session_state[unlock_key] = False
            st.rerun()
    else:
        with st.expander("🔒 下载原文件（需要密码）"):
            pwd = st.text_input(
                "下载密码", type="password", key=f"download_pwd_{book_id}")
            if st.button("验证密码并解锁下载", key=f"download_unlock_{book_id}"):
                if verify_download_password(pwd):
                    st.session_state[unlock_key] = True
                    st.rerun()
                else:
                    st.error("下载密码错误。")


@st.cache_data(show_spinner=False)
def load_disease_options():
    """“屡试屡验方”（排除“被忽略的名方”）中清理、去重后的病名菜单。"""
    conn = get_connection()
    try:
        df = pd.read_sql(
            "SELECT 临床病名 FROM formulas WHERE TRIM(方别) = ?",
            conn,
            params=("屡试屡验方",),
        )
    finally:
        conn.close()
    names = {normalize_disease_name(x) for x in df["临床病名"].tolist()}
    names.discard("")
    return sorted(names)


@st.cache_data(show_spinner=False)
def load_disease_index():
    """病名 -> 记录 ID 的映射，供“方式二”精确检索。"""
    conn = get_connection()
    try:
        df = pd.read_sql("SELECT id, 临床病名 FROM formulas", conn)
    finally:
        conn.close()
    mapping = {}
    for record_id, raw in zip(df["id"].tolist(), df["临床病名"].tolist()):
        name = normalize_disease_name(raw)
        if name:
            mapping.setdefault(name, []).append(int(record_id))
    return mapping


@st.cache_data(show_spinner=False)
def load_verif_options():
    """各种验方表中去重后的主治疾病菜单。"""
    conn = get_connection()
    try:
        df = pd.read_sql(f'SELECT 主治疾病 FROM "{VERIF_TABLE}"', conn)
    finally:
        conn.close()
    names = set()
    for raw in df["主治疾病"].tolist():
        for part in split_verif_indications(raw):
            name = normalize_disease_name(part)
            if name:
                names.add(name)
    return sorted(names)


@st.cache_data(show_spinner=False)
def load_verif_index():
    """主治疾病 -> 记录 ID 的映射，供各种验方的“方式二”检索。"""
    conn = get_connection()
    try:
        df = pd.read_sql(f'SELECT id, 主治疾病 FROM "{VERIF_TABLE}"', conn)
    finally:
        conn.close()
    mapping = {}
    for record_id, raw in zip(df["id"].tolist(), df["主治疾病"].tolist()):
        for part in split_verif_indications(raw):
            name = normalize_disease_name(part)
            if name:
                ids = mapping.setdefault(name, [])
                if int(record_id) not in ids:
                    ids.append(int(record_id))
    return mapping


def clear_search_caches():
    st.cache_data.clear()
    st.session_state.pop("verif_search_state", None)
    st.session_state.pop("formula_search_state", None)
    for key in list(st.session_state.keys()):
        if key.startswith("pdf_doc_"):
            doc = st.session_state.pop(key, None)
            try:
                if doc is not None:
                    doc.close()
            except Exception:
                pass
        elif key.startswith(("reader_pages_", "pdf_bytes_", "pdf_page_cache_")):
            st.session_state.pop(key, None)


bootstrap_book_library()

# --- 2. 侧边栏与路由 ---
st.sidebar.markdown("<h1 style='text-align: center; font-size: 60px;'>☯️</h1>", unsafe_allow_html=True)
st.sidebar.title("📜 中医典藏系统V2.0")
page = st.sidebar.radio("模块导航", [
    "🔍 传世名方基础检索",
    "🤖 传世名方 AI 检索",
    "📚 图书阅读",
    "🧾 各种验方检索",
    "🍲 四季养生食谱检索",
    "⚙️ 典藏后台管理系统"
])

st.sidebar.markdown("""
<div class="disclaimer">
<b>⚠️ 免责声明：</b><br>
本系统纯属医学资料检索系统，不作辩证诊断，结果仅供参考，须在医生指导下使用，否则后果自负。
</div>
""", unsafe_allow_html=True)

# ==========================================
# 模块一：传世名方基础检索
# ==========================================
if page == "🔍 传世名方基础检索":
    st.title("🔍 传世名方基础检索")
    st.markdown("仅在数据表的 **【临床病名】** 列中检索。可输入疾病名称，也可从菜单选择疾病名称。")

    search_mode = st.radio(
        "检索方式",
        ["方式一：输入疾病名称检索", "方式二：选择疾病名称检索"],
        horizontal=True,
        key="formula_search_mode",
    )

    search_term = ""
    submitted = False

    if search_mode == "方式一：输入疾病名称检索":
        search_term = st.text_input(
            "💡 方式一：输入疾病名称检索（例如：感冒、十二指肠溃疡）",
            "",
            key="formula_text_term",
        )
        submitted = st.button("开始常规检索", type="primary", key="formula_text_search")
    else:
        disease_options = load_disease_options()
        if not disease_options:
            st.info("📉 暂无可选择的疾病名称，请先在后台导入“屡试屡验方”数据。")
        else:
            search_term = st.selectbox(
                "📋 方式二：选择疾病名称检索",
                disease_options,
                index=None,
                placeholder="请输入或选择疾病名称…",
                key="formula_select_term",
            )
            submitted = st.button("开始常规检索", type="primary", key="formula_select_search")

    if submitted:
        term = str(search_term).strip() if search_term else ""
        if not term:
            st.session_state.pop("formula_search_state", None)
            st.warning("请先输入或选择疾病名称。")
        else:
            log_query("基础名方检索", term)
            if search_mode == "方式一：输入疾病名称检索":
                conn = get_connection()
                try:
                    total = pd.read_sql(
                        "SELECT COUNT(*) AS c FROM formulas WHERE 临床病名 LIKE ?",
                        conn,
                        params=(f"%{term}%",),
                    ).iloc[0]["c"]
                finally:
                    conn.close()
                st.session_state["formula_search_state"] = {
                    "mode": search_mode,
                    "term": term,
                    "ids": None,
                    "total": int(total),
                    "page": 1,
                }
            else:
                ids = sorted(set(load_disease_index().get(term, [])))
                st.session_state["formula_search_state"] = {
                    "mode": search_mode,
                    "term": term,
                    "ids": ids,
                    "total": len(ids),
                    "page": 1,
                }

    state = st.session_state.get("formula_search_state")
    if state:
        total = int(state.get("total", 0))
        total_pages = max(1, (total + FORMULA_PAGE_SIZE - 1) // FORMULA_PAGE_SIZE)
        page = max(1, min(int(state.get("page", 1)), total_pages))
        state["page"] = page

        st.caption(f"当前检索：{state.get('mode')} · {state.get('term')}")

        if total == 0:
            st.info("📉 浩瀚医海，未寻得完全匹配的记载。")
        else:
            offset = (page - 1) * FORMULA_PAGE_SIZE
            conn = get_connection()
            try:
                if state.get("mode") == "方式一：输入疾病名称检索":
                    # 只检索“临床病名”列，不再匹配方剂名、病例详情等其它列
                    df_result = pd.read_sql(
                        "SELECT * FROM formulas WHERE 临床病名 LIKE ? "
                        "ORDER BY id LIMIT ? OFFSET ?",
                        conn,
                        params=(f"%{state.get('term', '')}%", FORMULA_PAGE_SIZE, offset),
                    )
                else:
                    page_ids = (state.get("ids") or [])[offset:offset + FORMULA_PAGE_SIZE]
                    if page_ids:
                        placeholders = ",".join("?" * len(page_ids))
                        df_result = pd.read_sql(
                            f"SELECT * FROM formulas WHERE id IN ({placeholders}) ORDER BY id",
                            conn,
                            params=tuple(page_ids),
                        )
                    else:
                        df_result = pd.DataFrame()
            finally:
                conn.close()

            st.success(
                f"✅ 共检索到 {total} 则良方，当前第 {page}/{total_pages} 页（每页 {FORMULA_PAGE_SIZE} 条）：")
            for _, row in df_result.iterrows():
                render_formula_card(row)

            if total_pages > 1:
                prev_col, info_col, next_col = st.columns([1, 2, 1])
                with prev_col:
                    if st.button("⬅️ 上一页", disabled=page <= 1, key="formula_prev_page"):
                        state["page"] = page - 1
                        st.session_state["formula_search_state"] = state
                        st.rerun()
                with info_col:
                    st.markdown(
                        f"<div style='text-align:center;padding-top:6px'>第 {page} / {total_pages} 页</div>",
                        unsafe_allow_html=True)
                with next_col:
                    if st.button("下一页 ➡️", disabled=page >= total_pages, key="formula_next_page"):
                        state["page"] = page + 1
                        st.session_state["formula_search_state"] = state
                        st.rerun()

# ==========================================
# 模块二：传世名方 AI 智能检索
# ==========================================
elif page == "🤖 传世名方 AI 检索":
    st.title("🤖 AI 智能检索")
    st.markdown(
        "基于大模型的深度语义解析。**声明：AI 严格限定于本地《传世名方》数据库内进行比对筛查，宁缺毋滥，拒绝凭空伪造。**")

    if "ai_user_config" not in st.session_state:
        st.session_state.ai_user_config = dict(DEFAULT_AI_CONFIG)
    if "ai_storage_loaded" not in st.session_state:
        st.session_state.ai_storage_loaded = False
    if "ai_storage_action_id" not in st.session_state:
        st.session_state.ai_storage_action_id = 0

    action = st.session_state.pop("ai_storage_action", "")
    payload = st.session_state.pop("ai_storage_payload", None)
    action_id = int(st.session_state.get("ai_storage_action_id", 0))

    storage_component = _get_ai_storage_component()
    storage_result = None
    component_error = ""
    if storage_component is not None:
        try:
            storage_result = storage_component(
                action=action,
                payload=payload,
                action_id=action_id,
                key="ai_storage_bridge",
                default=None,
            )
        except Exception as e:
            component_error = str(e)
            storage_component = None

    if action:
        if storage_component is None:
            st.session_state["ai_settings_notice"] = (
                "warning",
                f"⚠️ 浏览器存储组件不可用，设置仅在本次会话有效。{component_error}".strip(),
            )
        elif action == "save":
            st.session_state["ai_settings_notice"] = (
                "success",
                "✅ AI 设置已保存到当前浏览器，刷新或重新打开页面后仍会保留。",
            )
        else:
            st.session_state["ai_settings_notice"] = (
                "success",
                "✅ 已清空当前浏览器中保存的 AI 设置。",
            )
    elif not st.session_state.ai_storage_loaded and isinstance(storage_result, dict):
        loaded_settings = storage_result.get("settings")
        if isinstance(loaded_settings, dict):
            loaded_config = dict(DEFAULT_AI_CONFIG)
            for key in loaded_config:
                value = loaded_settings.get(key)
                if isinstance(value, str):
                    loaded_config[key] = value.strip()
            st.session_state.ai_user_config = loaded_config
            st.session_state["ai_user_api_key"] = loaded_config["api_key"]
            st.session_state["ai_user_base_url"] = loaded_config["base_url"]
            st.session_state["ai_user_model_name"] = loaded_config["model_name"]
        st.session_state.ai_storage_loaded = True

    notice = st.session_state.get("ai_settings_notice")
    if notice:
        notice_kind, notice_text = notice
        getattr(st, notice_kind)(notice_text)

    saved_config = st.session_state.ai_user_config
    with st.expander("⚙️ AI 引擎设置（用户设置）", expanded=not bool(saved_config.get("api_key"))):
        st.caption("设置保存在当前浏览器的本地存储（localStorage）中，刷新或关闭后仍会保留；仅对当前浏览器和当前访问地址有效。可随时修改、保存或清空。")
        for widget_key, widget_value in (
            ("ai_user_api_key", saved_config.get("api_key", "")),
            ("ai_user_base_url", saved_config.get("base_url", DEFAULT_AI_CONFIG["base_url"])),
            ("ai_user_model_name", saved_config.get("model_name", DEFAULT_AI_CONFIG["model_name"])),
        ):
            if widget_key not in st.session_state:
                st.session_state[widget_key] = widget_value
        api_key = st.text_input("🔑 我的 API Key", type="password", key="ai_user_api_key",
                                on_change=_dismiss_ai_settings_notice)
        base_url = st.text_input("🌐 Base URL (接口地址)", key="ai_user_base_url",
                                 on_change=_dismiss_ai_settings_notice)
        model_name = st.text_input("🧠 调用的模型名称 (Model Name)", key="ai_user_model_name",
                                   on_change=_dismiss_ai_settings_notice)
        save_col, clear_col = st.columns(2)
        with save_col:
            st.button("💾 保存我的 AI 设置", key="ai_user_save", on_click=_save_ai_settings_callback)
        with clear_col:
            st.button("🗑️ 清空我的 AI 设置", key="ai_user_clear", on_click=_clear_ai_settings_callback)

    condition_desc = st.text_area(
        "✍️ 请详细描述疾病情况 (如: 胃痛时而剧烈，隐隐作痛，伴有头晕不眠，舌淡白...)",
        height=150,
        key="ai_condition_desc",
    )

    if st.button("🧠 提交 AI 深度匹配", type="primary", key="ai_submit"):
        if not condition_desc.strip():
            st.warning("请详细描述病症！")
        elif not api_key or not api_key.strip():
            st.error("⚠️ AI 尚未配置，请先在上方【AI 引擎设置（用户设置）】中填写您自己的 API Key。")
        else:
            with st.spinner("AI 正在典藏库中为您彻查比对...这可能需要十几秒，请稍候。"):
                conn = get_connection()
                try:
                    df_all = pd.read_sql(
                        "SELECT id, 方剂名, 临床病名, 方剂详情, 病例详情 FROM formulas", conn)
                finally:
                    conn.close()

                corpus = df_all.to_dict(orient="records")
                corpus_str = json.dumps(corpus[:100], ensure_ascii=False)

                system_prompt = """你是一位资深严谨的中医专家，请根据你所掌握的中医知识和经验，针对用户的【病情描述】进行全面的综合分析或将【病情描述】切成分词进行文字匹配，严格按照以下规则在我提供的<数据库语料>中推荐合适的方剂：
                    规则1：你只能使用我提供的<数据库语料>来分析或文字匹配检索，绝对不能编造外部知识！如果确实没有合适的，请返回空列表。
                    规则2：根据用户的【病情描述】，在语料中全面筛查分析，或将【病情描述】切成分词，在我提供的<数据库语料>中进行文字匹配，选择合适或匹配的方剂，按照合适度或匹配度减序排序。
                    规则3：输出严格的JSON。格式: {"matched_ids": [1, 5], "highlight_keywords": ["胃痛", "舌淡白"]}"""

                user_prompt = f"<数据库语料>\n{corpus_str}\n</数据库语料>\n\n用户的病情：{condition_desc}"

                try:
                    # 延迟导入：即使 AI 依赖未安装或版本异常，基础检索模块仍可正常使用
                    from openai import OpenAI

                    client = OpenAI(
                        api_key=api_key.strip(),
                        base_url=base_url.strip() or None,
                    )
                    response = client.chat.completions.create(
                        model=model_name.strip() or "deepseek-chat",
                        messages=[{"role": "system", "content": system_prompt},
                                  {"role": "user", "content": user_prompt}],
                        temperature=0.1,
                        response_format={"type": "json_object"}
                    )
                    ai_result = json.loads(response.choices[0].message.content)
                    matched_ids = ai_result.get("matched_ids", [])
                    keywords = ai_result.get("highlight_keywords", [])
                    log_query("AI深度检索", condition_desc)

                    if not matched_ids:
                        st.info("🤖 经过系统缜密分析，本地传世名方库中暂无与此高度契合的方剂。")
                    else:
                        st.success(f"🤖 已为您筛出最契合的 {len(matched_ids)} 剂良方（高亮显示病机词）：")
                        conn = get_connection()
                        try:
                            placeholders = ",".join("?" * len(matched_ids))
                            df_matched = pd.read_sql(
                                f"SELECT * FROM formulas WHERE id IN ({placeholders})",
                                conn,
                                params=tuple(matched_ids),
                            )
                        finally:
                            conn.close()

                        df_matched['id_cat'] = pd.Categorical(
                            df_matched['id'], categories=matched_ids, ordered=True)
                        df_matched = df_matched.sort_values('id_cat')

                        for _, row in df_matched.iterrows():
                            with st.container(border=True):
                                st.markdown(f"### 🛡️ AI 推荐：【{safe_text(row.get('方剂名'), '未知')}】")
                                st.markdown(
                                    f"**👨‍⚕️ 医家:** {safe_text(row.get('大医名'))} &nbsp;|&nbsp; "
                                    f"**🏷️ 方别:** {safe_text(row.get('方别'))} &nbsp;|&nbsp; "
                                    f"**🩺 适用病名:** {safe_text(row.get('临床病名'))}")
                                st.markdown(
                                    f"**📜 方剂详情:**<br>{highlight_text(str(row.get('方剂详情', '')), keywords)}",
                                    unsafe_allow_html=True)
                                st.markdown(
                                    f"**📝 临证提要:**<br>{highlight_text(str(row.get('临证提要', '')), keywords)}",
                                    unsafe_allow_html=True)
                                st.markdown(
                                    f"**📖 病例详情:**<br>{highlight_text(str(row.get('病例详情', '')), keywords)}",
                                    unsafe_allow_html=True)
                                if safe_text(row.get('按语')):
                                    st.markdown(
                                        f"**💡 按语:**<br>{highlight_text(str(row.get('按语')), keywords)}",
                                        unsafe_allow_html=True)
                                if safe_text(row.get('现代研究')):
                                    st.markdown(
                                        f"**🔬 现代研究:**<br>{highlight_text(str(row.get('现代研究')), keywords)}",
                                        unsafe_allow_html=True)
                except Exception as e:
                    st.error(f"AI 调用失败，请检查配置与网络：{e}")

# ==========================================
# 模块三：图书阅读
# ==========================================
elif page == "📚 图书阅读":
    st.title("📚 图书阅读")
    st.markdown("按分类选择书名阅读。阅读位置会自动记忆，下次打开同一本书会回到上次位置。")

    reader_category = st.selectbox("图书分类", BOOK_CATEGORIES, key="reader_category")
    books = load_books(reader_category)

    if books.empty:
        st.info("📭 该分类暂无图书。请进入【典藏后台管理系统 → 📚 图书库】批量导入。")
    else:
        selected_title = st.selectbox(
            "选择书名", books["title"].tolist(),
            index=None,
            placeholder="请输入或选择书名…",
            key=f"reader_book_{reader_category}")
        if not selected_title:
            st.info("请从上方选择书名开始阅读。")
            st.stop()
        book_row = books[books["title"] == selected_title].iloc[0]
        book_id = int(book_row["id"])
        book = load_book_text(book_id)

        if not book:
            st.error("图书记录不存在，请刷新页面或重新导入。")
        else:
            title, fmt, text, char_count, page_count, file_size = book
            st.subheader(f"📖 {safe_text(title)}")
            st.caption(
                f"{reader_category} · {safe_text(fmt).upper()} · "
                f"{format_file_size(file_size)} · {int(char_count or 0)} 字")
            text = text or ""

            if text.strip():
                font_size = st.slider(
                    "字体大小（文本阅读）", 14, 30, 18, 1, key="reader_font_size")
                pages_key = f"reader_pages_{book_id}"
                length_key = f"{pages_key}_length"
                if st.session_state.get(length_key) != len(text):
                    st.session_state[pages_key] = build_book_pages(text)
                    st.session_state[length_key] = len(text)
                pages = st.session_state.get(pages_key) or []
                if pages:
                    position = get_reading_position(book_id)
                    position = max(0, min(position, len(text) - 1))
                    starts = [p[0] for p in pages]
                    page_index = bisect.bisect_right(starts, position) - 1
                    page_index = max(0, min(page_index, len(pages) - 1))
                    if position != starts[page_index]:
                        save_reading_position(book_id, starts[page_index])

                    page_slider_key = f"text_page_slider_{book_id}"
                    force_key = f"text_force_page_{book_id}"
                    force_page = st.session_state.pop(force_key, False)
                    if force_page or page_slider_key not in st.session_state:
                        st.session_state[page_slider_key] = page_index + 1
                    selected_page = st.slider(
                        "滑动翻页", 1, len(pages), key=page_slider_key)
                    if selected_page != page_index + 1:
                        save_reading_position(book_id, pages[selected_page - 1][0])
                        st.session_state[force_key] = True
                        st.rerun()

                    st.caption("阅读位置会自动保存；可拖动滑块或点击按钮翻页。")
                    render_book_nav(
                        book_id, page_index, pages, "top", force_key)
                    st.markdown(
                        book_page_html(pages[page_index][1], font_size),
                        unsafe_allow_html=True)
                    render_book_nav(
                        book_id, page_index, pages, "bottom", force_key)
            elif safe_text(fmt) == "pdf":
                bytes_key = f"pdf_bytes_{book_id}"
                if bytes_key not in st.session_state:
                    st.session_state[bytes_key] = load_book_file(book_id)
                content = st.session_state.get(bytes_key)
                if content:
                    total_pages = int(page_count or 0) or get_pdf_page_count(content)
                    if total_pages <= 0:
                        st.warning("无法读取 PDF 页数，可下载原文件阅读。")
                    else:
                        page = max(1, min(get_reading_position(book_id) + 1, total_pages))
                        st.info(
                            "该 PDF 未提取到文字（可能是扫描版），按原版页面渲染阅读；"
                            "页码位置会自动记忆，可调节缩放。")
                        pdf_zoom = st.slider(
                            "PDF 缩放", 1.0, 3.0, 1.5, 0.1, key=f"pdf_zoom_{book_id}")
                        page_slider_key = f"pdf_page_slider_{book_id}"
                        force_key = f"pdf_force_page_{book_id}"
                        force_page = st.session_state.pop(force_key, False)
                        if force_page or page_slider_key not in st.session_state:
                            st.session_state[page_slider_key] = page
                        selected_page = st.slider(
                            "滑动翻页", 1, total_pages, key=page_slider_key)
                        if selected_page != page:
                            save_reading_position(book_id, selected_page - 1)
                            st.session_state[force_key] = True
                            st.rerun()

                        image_data = render_pdf_page_cached(book_id, content, page - 1)
                        if image_data:
                            image_b64 = base64.b64encode(image_data).decode("ascii")
                            zoom_percent = int(round(pdf_zoom * 100))
                            st.markdown(
                                '<div class="page-fade" style="overflow:auto;max-height:82vh;'
                                'border:1px solid #E8DCC4;border-radius:8px;'
                                'background:#FFFFFF;padding:8px;">'
                                f'<img src="data:image/jpeg;base64,{image_b64}" '
                                f'style="width:{zoom_percent}%;max-width:none;display:block;'
                                'margin:0 auto;" />'
                                '</div>',
                                unsafe_allow_html=True)
                            st.caption(f"当前缩放：{zoom_percent}%")
                        else:
                            st.warning("PDF 页面渲染失败，可下载原文件阅读。")

                        prev_col, info_col, next_col = st.columns([1, 2, 1])
                        with prev_col:
                            if st.button("⬅️ 上一页", disabled=page <= 1,
                                         key=f"pdf_prev_{book_id}"):
                                save_reading_position(book_id, page - 2)
                                st.session_state[force_key] = True
                                st.rerun()
                        with info_col:
                            st.markdown(
                                f"<div style='text-align:center;padding-top:6px'>"
                                f"第 {page} / {total_pages} 页</div>",
                                unsafe_allow_html=True)
                        with next_col:
                            if st.button("下一页 ➡️", disabled=page >= total_pages,
                                         key=f"pdf_next_{book_id}"):
                                save_reading_position(book_id, page)
                                st.session_state[force_key] = True
                                st.rerun()
                    download_name = safe_text(book_row.get("filename")) or f"{title}.pdf"
                    render_protected_download(
                        book_id, download_name, content, "application/pdf")
                else:
                    st.warning("未找到 PDF 原文件。")
            else:
                st.warning(
                    "该文件未能提取出可阅读文字，可能是扫描版或旧版 DOC/WPS 格式。"
                    "可以下载原文件阅读。")
                content = load_book_file(book_id)
                if content:
                    download_name = safe_text(book_row.get("filename")) or f"{title}.{fmt}"
                    render_protected_download(
                        book_id, download_name, content, "application/octet-stream")

# ==========================================
# 模块四：各种验方检索
# ==========================================
elif page == "🧾 各种验方检索":
    st.title("🧾 各种验方检索")
    st.markdown("在 **【各种验方表】** 中检索。表格字段：方名、方剂及用法、主治疾病、症状、案例、出处。支持输入疾病名称或症状，也支持从菜单中选择疾病名称（按主治疾病整理）。")

    conn = get_connection()
    try:
        # 防止热重载时启动迁移未执行，进入页面时再确保表结构正确
        migrate_verif_table(conn)
        conn.commit()
        verif_count = pd.read_sql(f'SELECT COUNT(*) AS c FROM "{VERIF_TABLE}"', conn).iloc[0]["c"]
    finally:
        conn.close()

    if not verif_count:
        st.info("📭 【各种验方表】暂无数据。请进入【典藏后台管理系统 → 📁 批量导入】导入数据后再检索。")

    search_mode = st.radio(
        "检索方式",
        ["方式一：输入疾病名称检索", "方式二：选择疾病名称检索"],
        horizontal=True,
        key="verif_search_mode",
    )

    search_term = ""
    submitted = False

    if search_mode == "方式一：输入疾病名称检索":
        search_term = st.text_input(
            "💡 方式一：输入疾病名称检索（例如：感冒、咳嗽，也可输入方名或症状）",
            "",
            key="verif_text_term",
        )
        submitted = st.button("开始验方检索", type="primary", key="verif_text_search")
    else:
        verif_options = load_verif_options()
        if not verif_options:
            st.info("📉 暂无可选择的疾病名称，请先在后台导入“各种验方表”数据。")
        else:
            search_term = st.selectbox(
                "📋 方式二：选择疾病名称检索",
                verif_options,
                index=None,
                placeholder="请输入或选择疾病名称…",
                key="verif_select_term",
            )
            submitted = st.button("开始验方检索", type="primary", key="verif_select_search")

    if submitted:
        term = str(search_term).strip() if search_term else ""
        if not term:
            st.session_state.pop("verif_search_state", None)
            st.warning("请先输入或选择疾病名称。")
        else:
            log_query("各种验方检索", term)
            if search_mode == "方式一：输入疾病名称检索":
                conn = get_connection()
                try:
                    like = f"%{term}%"
                    total = pd.read_sql(
                        f'SELECT COUNT(*) AS c FROM "{VERIF_TABLE}" '
                        "WHERE 主治疾病 LIKE ? OR 方名 LIKE ? OR 症状 LIKE ?",
                        conn,
                        params=(like, like, like),
                    ).iloc[0]["c"]
                finally:
                    conn.close()
                st.session_state["verif_search_state"] = {
                    "mode": search_mode,
                    "term": term,
                    "ids": None,
                    "total": int(total),
                    "page": 1,
                }
            else:
                ids = sorted(set(load_verif_index().get(term, [])))
                st.session_state["verif_search_state"] = {
                    "mode": search_mode,
                    "term": term,
                    "ids": ids,
                    "total": len(ids),
                    "page": 1,
                }

    state = st.session_state.get("verif_search_state")
    if state:
        total = int(state.get("total", 0))
        total_pages = max(1, (total + VERIF_PAGE_SIZE - 1) // VERIF_PAGE_SIZE)
        page = max(1, min(int(state.get("page", 1)), total_pages))
        state["page"] = page

        st.caption(f"当前检索：{state.get('mode')} · {state.get('term')}")

        if total == 0:
            st.info("📉 未寻得匹配的验方记载。")
        else:
            offset = (page - 1) * VERIF_PAGE_SIZE
            conn = get_connection()
            try:
                if state.get("mode") == "方式一：输入疾病名称检索":
                    like = f"%{state.get('term', '')}%"
                    df_result = pd.read_sql(
                        f'SELECT * FROM "{VERIF_TABLE}" '
                        "WHERE 主治疾病 LIKE ? OR 方名 LIKE ? OR 症状 LIKE ? "
                        "ORDER BY id LIMIT ? OFFSET ?",
                        conn,
                        params=(like, like, like, VERIF_PAGE_SIZE, offset),
                    )
                else:
                    page_ids = (state.get("ids") or [])[offset:offset + VERIF_PAGE_SIZE]
                    if page_ids:
                        placeholders = ",".join("?" * len(page_ids))
                        df_result = pd.read_sql(
                            f'SELECT * FROM "{VERIF_TABLE}" WHERE id IN ({placeholders}) ORDER BY id',
                            conn,
                            params=tuple(page_ids),
                        )
                    else:
                        df_result = pd.DataFrame()
            finally:
                conn.close()

            st.success(
                f"✅ 共检索到 {total} 则验方，当前第 {page}/{total_pages} 页（每页 {VERIF_PAGE_SIZE} 条）：")
            for _, row in df_result.iterrows():
                render_verif_card(row)

            if total_pages > 1:
                prev_col, info_col, next_col = st.columns([1, 2, 1])
                with prev_col:
                    if st.button("⬅️ 上一页", disabled=page <= 1, key="verif_prev_page"):
                        state["page"] = page - 1
                        st.session_state["verif_search_state"] = state
                        st.rerun()
                with info_col:
                    st.markdown(
                        f"<div style='text-align:center;padding-top:6px'>第 {page} / {total_pages} 页</div>",
                        unsafe_allow_html=True)
                with next_col:
                    if st.button("下一页 ➡️", disabled=page >= total_pages, key="verif_next_page"):
                        state["page"] = page + 1
                        st.session_state["verif_search_state"] = state
                        st.rerun()

# ==========================================
# 模块四：养生食谱检索
# ==========================================
elif page == "🍲 四季养生食谱检索":
    st.title("🍲 养生食谱数据库")
    st.markdown("按食材用料、核心功效、适用季节、菜式品种四个维度检索。不填的条件默认不作限制。")

    with st.container(border=True):
        col1, col2 = st.columns(2)
        with col1:
            cailiao = st.text_input("🥦 食材用料 (如: 猪肉、山药)", key="recipe_cailiao")
            gongxiao = st.text_input("✨ 核心功效 (如: 健脾、润肺)", key="recipe_gongxiao")
        with col2:
            jijie = st.text_input("🌸 适用季节 (如: 春季、夏季)", key="recipe_jijie")
            pinlei = st.text_input("🥘 菜式品种 (如: 汤羹、粥)", key="recipe_pinlei")

        if st.button("查询养生食谱", type="primary", key="recipe_search"):
            query = "SELECT * FROM recipes WHERE 1=1"
            params = []
            if cailiao:
                query += " AND 材料和制作 LIKE ?"
                params.append(f'%{cailiao}%')
            if gongxiao:
                query += " AND 功效 LIKE ?"
                params.append(f'%{gongxiao}%')
            if jijie:
                query += " AND 适合季节 LIKE ?"
                params.append(f'%{jijie}%')
            if pinlei:
                query += " AND 品类 LIKE ?"
                params.append(f'%{pinlei}%')
            query += " LIMIT 150"  # 限制150条防卡顿

            conn = get_connection()
            try:
                df_result = pd.read_sql(query, conn, params=tuple(params))
            finally:
                conn.close()

            log_query("食谱检索", f"{cailiao}|{gongxiao}|{jijie}|{pinlei}")

            if df_result.empty:
                st.info("📉 未寻得符合全部条件的食谱，请尝试精简检索词。")
            else:
                st.success(f"✅ 寻得 {len(df_result)} 道养生佳肴：")
                for _, row in df_result.iterrows():
                    with st.expander(f"🍲 【{safe_text(row.get('菜名'), '未知')}】 | 类: {safe_text(row.get('品类'))}"):
                        # 逐列展示该结果的全部字段，不截断、不遗漏
                        for column in df_result.columns:
                            value = safe_text(row.get(column))
                            st.markdown(f"**{column}:**\n {value if value else '（空）'}")

# ==========================================
# 模块五：后台管理系统
# ==========================================
elif page == "⚙️ 典藏后台管理系统":
    st.title("⚙️ 系统中枢及数据管理")

    if "admin_auth" not in st.session_state:
        st.session_state.admin_auth = False

    if not st.session_state.admin_auth:
        pwd = st.text_input("🔑 请输入系统超级管理员密码", type="password", key="admin_pwd")
        if st.button("进入中枢", key="admin_login"):
            if verify_admin_password(pwd):
                st.session_state.admin_auth = True
                st.rerun()
            else:
                st.error("密码错误！")
    else:
        st.success("✅ 身份验证通过：田 夫 (为了系统流畅，请规范操作)")

        tab_stat, tab_edit, tab_import, tab_books, tab_password, tab_danger = st.tabs([
            "📊 数据统计", "✏️ 数据维护(增删)", "📁 批量导入",
            "📚 图书库", "🔐 密码管理", "⚠️ 高危操作"
        ])

        with tab_stat:
            st.subheader("平台数据态势")
            conn = get_connection()
            try:
                num_formulas = pd.read_sql("SELECT COUNT(*) as c FROM formulas", conn).iloc[0]['c']
                num_recipes = pd.read_sql("SELECT COUNT(*) as c FROM recipes", conn).iloc[0]['c']
                num_verif = pd.read_sql(f'SELECT COUNT(*) as c FROM "{VERIF_TABLE}"', conn).iloc[0]['c']
                num_books = pd.read_sql("SELECT COUNT(*) as c FROM books", conn).iloc[0]['c']
                num_logs = pd.read_sql("SELECT COUNT(*) as c FROM query_logs", conn).iloc[0]['c']
            finally:
                conn.close()
            sc1, sc2, sc3, sc4, sc5 = st.columns(5)
            sc1.metric("📚 传世名方总数", num_formulas)
            sc2.metric("🧾 各种验方总数", num_verif)
            sc3.metric("🍲 养生食谱总数", num_recipes)
            sc4.metric("📖 图书总数", num_books)
            sc5.metric("👣 历史总检索次数", num_logs)

            st.markdown("### 📝 最近访客检索日志")
            conn = get_connection()
            try:
                logs_df = pd.read_sql(
                    "SELECT id as 序号, query_time as 检索时间, search_type as 检索模块, keyword as 检索词 "
                    "FROM query_logs ORDER BY id DESC LIMIT 50",
                    conn)
            finally:
                conn.close()
            st.dataframe(logs_df)

        with tab_edit:
            st.subheader("✏️ 数据精确定向维护")
            st.warning("⚠️ 数据库体积庞大，全量铺开会导致浏览器崩溃卡死。请通过下方表单精确完成录入、修改与删除。")

            action_type = st.radio(
                "请选择执行的操作",
                ["➕ 新增记录", "🔍 搜索与修改", "🗑️ 定向删除"],
                horizontal=True,
                key="admin_action",
            )

            if action_type == "➕ 新增记录":
                tbl = st.selectbox(
                    "目标表",
                    ["formulas (名方)", "recipes (食谱)", "各种验方表 (验方)"],
                    key="admin_insert_table",
                )
                with st.form("insert_f"):
                    if "formulas" in tbl:
                        st.info("正在录入：传世名方")
                        f_name = st.text_input("方剂名*")
                        b_name = st.text_input("临床病名*")
                        d_name = st.text_input("大医名")
                        det = st.text_area("方剂详情")
                        cas = st.text_area("病例详情")
                        sumry = st.text_area("临证提要")
                        if st.form_submit_button("保存至名方库"):
                            if f_name and b_name:
                                conn = get_connection()
                                conn.execute(
                                    "INSERT INTO formulas (方剂名, 临床病名, 大医名, 方剂详情, 病例详情, 临证提要) VALUES (?,?,?,?,?,?)",
                                    (f_name, b_name, d_name, det, cas, sumry))
                                conn.commit()
                                conn.close()
                                clear_search_caches()
                                st.success("新增成功！")
                            else:
                                st.error("必填项不可为空！")
                    elif "recipes" in tbl:
                        st.info("正在录入：养生食谱")
                        r_name = st.text_input("菜名*")
                        r_type = st.text_input("品类")
                        r_mat = st.text_area("材料和制作")
                        r_eff = st.text_area("功效")
                        r_note = st.text_area("注意事项")
                        if st.form_submit_button("保存至食谱库"):
                            if r_name:
                                conn = get_connection()
                                conn.execute(
                                    "INSERT INTO recipes (菜名, 品类, 材料和制作, 功效, 注意事项) VALUES (?,?,?,?,?)",
                                    (r_name, r_type, r_mat, r_eff, r_note))
                                conn.commit()
                                conn.close()
                                st.success("新增成功！")
                            else:
                                st.error("菜名为必填项！")
                    else:
                        st.info("正在录入：各种验方")
                        v_name = st.text_input("方名*")
                        v_recipe = st.text_area("方剂及用法")
                        v_disease = st.text_area("主治疾病")
                        v_symptom = st.text_area("症状")
                        v_case = st.text_area("案例")
                        v_source = st.text_input("出处")
                        if st.form_submit_button("保存至验方库"):
                            if v_name:
                                conn = get_connection()
                                conn.execute(
                                    f'INSERT INTO "{VERIF_TABLE}" (方名, 方剂及用法, 主治疾病, 症状, 案例, 出处) '
                                    "VALUES (?,?,?,?,?,?)",
                                    (v_name, v_recipe, v_disease, v_symptom, v_case, v_source))
                                conn.commit()
                                conn.close()
                                clear_search_caches()
                                st.success("新增成功！")
                            else:
                                st.error("方名为必填项！")

            elif action_type == "🔍 搜索与修改":
                st.markdown("通过关键字定位数据并进行编辑更新。")
                edit_tbl = st.selectbox(
                    "修改目标表",
                    ["formulas (名方)", "recipes (食谱)", "各种验方表 (验方)"],
                    key="admin_edit_table",
                )
                if "formulas" in edit_tbl:
                    real_tbl = FORMULA_TABLE
                elif "recipes" in edit_tbl:
                    real_tbl = RECIPE_TABLE
                else:
                    real_tbl = VERIF_TABLE

                search_kw = st.text_input("🔍 输入检索词 (如方名、菜名、病名等以查找)", key="admin_edit_kw")

                if search_kw.strip():
                    conn = get_connection()
                    try:
                        if real_tbl == FORMULA_TABLE:
                            query = ("SELECT id, 方剂名, 临床病名, 大医名 FROM formulas "
                                     "WHERE 方剂名 LIKE ? OR 临床病名 LIKE ? LIMIT 20")
                        elif real_tbl == RECIPE_TABLE:
                            query = ("SELECT id, 菜名, 品类 FROM recipes "
                                     "WHERE 菜名 LIKE ? OR 功效 LIKE ? OR 注意事项 LIKE ? LIMIT 20")
                        else:
                            query = (f'SELECT id, 方名, 主治疾病 FROM "{VERIF_TABLE}" '
                                     "WHERE 方名 LIKE ? OR 主治疾病 LIKE ? OR 症状 LIKE ? LIMIT 20")
                        search_res = pd.read_sql(
                            query, conn,
                            params=tuple([f'%{search_kw}%'] * query.count("?")))
                    finally:
                        conn.close()

                    if search_res.empty:
                        st.info("未找到匹配的记录。")
                    else:
                        st.write(f"找到 {len(search_res)} 条关联记录 (最多显示前20条)：")
                        st.dataframe(search_res, hide_index=True)

                        edit_id = st.number_input("请输入您要修改的记录 ID", min_value=0, step=1, value=0)

                        if edit_id > 0:
                            conn = get_connection()
                            try:
                                record = pd.read_sql(
                                    f'SELECT * FROM "{real_tbl}" WHERE id=?', conn, params=(edit_id,))
                            finally:
                                conn.close()

                            if not record.empty:
                                row = record.iloc[0]
                                with st.form("update_form"):
                                    if real_tbl == FORMULA_TABLE:
                                        st.info(f"正在修改 ID: {edit_id} 的名方")
                                        u_f_name = st.text_input("方剂名*", value=safe_text(row.get('方剂名')))
                                        u_b_name = st.text_input("临床病名*", value=safe_text(row.get('临床病名')))
                                        u_d_name = st.text_input("大医名", value=safe_text(row.get('大医名')))
                                        u_det = st.text_area("方剂详情", value=safe_text(row.get('方剂详情')))
                                        u_cas = st.text_area("病例详情", value=safe_text(row.get('病例详情')))
                                        u_sumry = st.text_area("临证提要", value=safe_text(row.get('临证提要')))

                                        if st.form_submit_button("💾 保存名方修改", type="primary"):
                                            if u_f_name and u_b_name:
                                                conn = get_connection()
                                                conn.execute(
                                                    "UPDATE formulas SET 方剂名=?, 临床病名=?, 大医名=?, 方剂详情=?, 病例详情=?, 临证提要=? WHERE id=?",
                                                    (u_f_name, u_b_name, u_d_name, u_det, u_cas, u_sumry, edit_id))
                                                conn.commit()
                                                conn.close()
                                                clear_search_caches()
                                                st.success("名方记录修改成功！")
                                            else:
                                                st.error("必填项不可为空！")
                                    elif real_tbl == RECIPE_TABLE:
                                        st.info(f"正在修改 ID: {edit_id} 的食谱")
                                        u_r_name = st.text_input("菜名*", value=safe_text(row.get('菜名')))
                                        u_r_type = st.text_input("品类", value=safe_text(row.get('品类')))
                                        u_r_mat = st.text_area("材料和制作", value=safe_text(row.get('材料和制作')))
                                        u_r_eff = st.text_area("功效", value=safe_text(row.get('功效')))
                                        u_r_note = st.text_area("注意事项", value=safe_text(row.get('注意事项')))

                                        if st.form_submit_button("💾 保存食谱修改", type="primary"):
                                            if u_r_name:
                                                conn = get_connection()
                                                conn.execute(
                                                    "UPDATE recipes SET 菜名=?, 品类=?, 材料和制作=?, 功效=?, 注意事项=? WHERE id=?",
                                                    (u_r_name, u_r_type, u_r_mat, u_r_eff, u_r_note, edit_id))
                                                conn.commit()
                                                conn.close()
                                                st.success("食谱记录修改成功！")
                                            else:
                                                st.error("菜名为必填项！")
                                    else:
                                        st.info(f"正在修改 ID: {edit_id} 的验方")
                                        u_v_name = st.text_input("方名*", value=safe_text(row.get('方名')))
                                        u_v_recipe = st.text_area("方剂及用法", value=safe_text(row.get('方剂及用法')))
                                        u_v_disease = st.text_area("主治疾病", value=safe_text(row.get('主治疾病')))
                                        u_v_symptom = st.text_area("症状", value=safe_text(row.get('症状')))
                                        u_v_case = st.text_area("案例", value=safe_text(row.get('案例')))
                                        u_v_source = st.text_input("出处", value=safe_text(row.get('出处')))

                                        if st.form_submit_button("💾 保存验方修改", type="primary"):
                                            if u_v_name:
                                                conn = get_connection()
                                                conn.execute(
                                                    f'UPDATE "{VERIF_TABLE}" SET 方名=?, 方剂及用法=?, 主治疾病=?, 症状=?, 案例=?, 出处=? WHERE id=?',
                                                    (u_v_name, u_v_recipe, u_v_disease, u_v_symptom,
                                                     u_v_case, u_v_source, edit_id))
                                                conn.commit()
                                                conn.close()
                                                clear_search_caches()
                                                st.success("验方记录修改成功！")
                                            else:
                                                st.error("方名为必填项！")
                            else:
                                st.warning("未找到该 ID 对应的记录，请核对后输入。")

            elif action_type == "🗑️ 定向删除":
                st.markdown("通过 ID 精确定位数据并删除，确保系统响应极速不卡顿。")
                del_tbl = st.selectbox(
                    "目标表", [FORMULA_TABLE, RECIPE_TABLE, VERIF_TABLE], key="admin_del_table")
                del_id = st.number_input(
                    f"请输入要删除的 {del_tbl} 记录的 ID", min_value=1, step=1, key="admin_del_id")

                # 预览要删除的数据
                conn = get_connection()
                try:
                    preview = pd.read_sql(
                        f'SELECT * FROM "{del_tbl}" WHERE id=?', conn, params=(del_id,))
                finally:
                    conn.close()
                if not preview.empty:
                    st.dataframe(preview)
                    if st.button("🚨 确认删除此记录", key="admin_del_confirm"):
                        conn = get_connection()
                        conn.execute(f'DELETE FROM "{del_tbl}" WHERE id=?', (del_id,))
                        conn.commit()
                        conn.close()
                        clear_search_caches()
                        st.success("已成功删除！")
                        st.rerun()
                else:
                    st.info("未找到此 ID 对应的数据。")

        with tab_import:
            st.subheader("📥 Excel 批量智能导入")
            uploaded_file = st.file_uploader("请上传待导入的 Excel 文件 (.xlsx)", type=["xlsx", "xls"])
            target_label = st.radio(
                "导入目标库",
                ["formulas (名方)", "recipes (食谱)", "各种验方表 (验方)"],
                key="admin_import_table",
            )
            if "formulas" in target_label:
                target_table = FORMULA_TABLE
            elif "recipes" in target_label:
                target_table = RECIPE_TABLE
            else:
                target_table = VERIF_TABLE

            if target_table == VERIF_TABLE:
                st.caption("各种验方表模板列：方名、方剂及用法、主治疾病、症状、案例、出处。")

            if uploaded_file is not None:
                xls = pd.ExcelFile(uploaded_file)
                st.write(f"📂 检测到 {len(xls.sheet_names)} 个工作表")
                if st.button("开始批量读取并智能去重导入", type="primary", key="admin_import_run"):
                    try:
                        conn = get_connection()
                        if target_table == VERIF_TABLE:
                            # 兼容旧库/热重载未迁移的情况，导入前强制校正表结构
                            migrate_verif_table(conn)
                            conn.commit()
                            df_list = [pd.read_excel(xls, sheet_name=s) for s in xls.sheet_names]
                            new_data = pd.concat(df_list, ignore_index=True)
                            new_data.columns = [str(c).strip() for c in new_data.columns]
                            # 兼容旧版表头：方剂/主治功能/来源 -> 模板新列名
                            for old_col, new_col in (
                                ("方剂", "方剂及用法"),
                                ("主治功能", "主治疾病"),
                                ("来源", "出处"),
                            ):
                                if old_col in new_data.columns and new_col not in new_data.columns:
                                    new_data.rename(columns={old_col: new_col}, inplace=True)
                            # 模板没有“备注”列，旧数据中的备注并入案例，避免丢失
                            if "备注" in new_data.columns:
                                if "案例" not in new_data.columns:
                                    new_data["案例"] = ""
                                base_case = new_data["案例"].fillna("").astype(str)
                                note_text = new_data["备注"].fillna("").astype(str).str.strip()
                                note_mask = note_text != ""
                                new_data.loc[note_mask, "案例"] = (
                                    base_case[note_mask].where(
                                        base_case[note_mask] == "", base_case[note_mask] + "\n")
                                    + "【备注】" + note_text[note_mask])
                                new_data = new_data.drop(columns=["备注"])
                            for col in VERIF_COLUMNS:
                                if col not in new_data.columns:
                                    new_data[col] = ""
                            new_data = new_data[VERIF_COLUMNS].fillna("")

                            verif_col_sql = ", ".join(f'"{c}"' for c in VERIF_COLUMNS)
                            old_df = pd.read_sql(
                                f'SELECT {verif_col_sql} FROM "{VERIF_TABLE}"', conn)
                            existing = {
                                tuple(safe_text(r[c]) for c in VERIF_COLUMNS)
                                for _, r in old_df.iterrows()
                            }
                            rows_to_insert = []
                            for _, r in new_data.iterrows():
                                values = tuple(safe_text(r[c]) for c in VERIF_COLUMNS)
                                if not any(values) or values in existing:
                                    continue
                                rows_to_insert.append(values)
                                existing.add(values)
                            if rows_to_insert:
                                conn.executemany(
                                    f'INSERT INTO "{VERIF_TABLE}" '
                                    "(方名, 方剂及用法, 主治疾病, 症状, 案例, 出处) "
                                    "VALUES (?,?,?,?,?,?)",
                                    rows_to_insert)
                            inserted = len(rows_to_insert)
                            conn.commit()
                            total = conn.execute(
                                f'SELECT COUNT(*) FROM "{VERIF_TABLE}"').fetchone()[0]
                            conn.close()
                            clear_search_caches()
                            st.success(f"🎉 导入成功！新增 {inserted} 条，过滤重复后验方库中共存有 {total} 条数据。")
                        else:
                            old_df = pd.read_sql(f'SELECT * FROM "{target_table}"', conn)
                            df_list = [pd.read_excel(xls, sheet_name=s) for s in xls.sheet_names]
                            new_data = pd.concat(df_list, ignore_index=True)
                            new_data.columns = [str(c).strip() for c in new_data.columns]
                            # 忽略上传文件中的 id 列，新增记录由数据库统一分配主键
                            if "id" in new_data.columns:
                                new_data = new_data.drop(columns=["id"])

                            # 兼容旧版食谱表头：功效和注意事项 -> 功效；季节功能 -> 注意事项
                            if target_table == RECIPE_TABLE:
                                if "功效和注意事项" in new_data.columns:
                                    if "功效" not in new_data.columns:
                                        new_data.rename(columns={"功效和注意事项": "功效"}, inplace=True)
                                    else:
                                        new_data["功效"] = new_data["功效"].where(
                                            new_data["功效"].fillna("").astype(str).str.strip() != "",
                                            new_data["功效和注意事项"])
                                        new_data.drop(columns=["功效和注意事项"], inplace=True)
                                if "季节功能" in new_data.columns:
                                    if "注意事项" not in new_data.columns:
                                        new_data.rename(columns={"季节功能": "注意事项"}, inplace=True)
                                    else:
                                        new_data["注意事项"] = new_data["注意事项"].where(
                                            new_data["注意事项"].fillna("").astype(str).str.strip() != "",
                                            new_data["季节功能"])
                                        new_data.drop(columns=["季节功能"], inplace=True)

                            combined_df = pd.concat([old_df, new_data], ignore_index=True)

                            if target_table == FORMULA_TABLE and '方剂名' in combined_df.columns:
                                combined_df.drop_duplicates(subset=['方剂名', '临床病名'], keep='last', inplace=True)
                            elif target_table == RECIPE_TABLE and '菜名' in combined_df.columns:
                                combined_df.drop_duplicates(subset=['菜名'], keep='last', inplace=True)

                            combined_df.to_sql(target_table, conn, if_exists="replace", index=False)
                            rebuild_imported_table(conn, target_table, TABLE_COLUMNS[target_table])
                            conn.commit()
                            conn.close()
                            clear_search_caches()
                            st.success(f"🎉 导入成功！过滤重复后库中共存有 {len(combined_df)} 条数据。")
                    except Exception as e:
                        st.error(f"导入出错，可能是文件格式不匹配: {e}")

        with tab_books:
            st.subheader("📚 图书批量导入")
            st.caption("支持 PDF、DOC、DOCX、WPS、TXT、MD、RTF；书名取文件名（不含扩展名）。")
            book_category = st.selectbox(
                "图书分类", BOOK_CATEGORIES, key="book_import_category")
            uploaded_books = st.file_uploader(
                "选择图书文件（可多选）", type=BOOK_FILE_TYPES,
                accept_multiple_files=True, key="book_import_files")

            if st.button("开始导入图书", type="primary", key="book_import_run"):
                if not uploaded_books:
                    st.warning("请先选择图书文件。")
                else:
                    conn = get_connection()
                    success_count = 0
                    warnings = []
                    try:
                        total = len(uploaded_books)
                        progress = st.progress(0.0, text="正在导入图书...")
                        for i, uploaded in enumerate(uploaded_books, start=1):
                            filename = os.path.basename(uploaded.name)
                            title = os.path.splitext(filename)[0].strip() or "未命名图书"
                            fmt = os.path.splitext(filename)[1].lower().lstrip(".")
                            data = uploaded.getvalue()
                            try:
                                text, page_count, error = extract_book_content(data, filename)
                                upsert_book(
                                    conn, book_category, title, fmt, filename,
                                    data, text, page_count)
                                conn.commit()
                                success_count += 1
                                if error:
                                    warnings.append(f"{filename}：{error}")
                            except Exception as e:
                                warnings.append(f"{filename}：导入失败 {e}")
                            progress.progress(i / total, text=f"已处理 {i}/{total}：{filename}")
                        progress.progress(1.0, text="图书导入完成")
                    finally:
                        conn.close()
                    clear_search_caches()
                    backup_book_library()
                    st.success(f"✅ 图书导入完成，共处理 {success_count} 个文件。")
                    for message in warnings:
                        st.warning(message)

            st.markdown("### 现有图书")
            conn = get_connection()
            try:
                books_df = pd.read_sql(
                    "SELECT id, category AS 分类, title AS 书名, format AS 格式, "
                    "file_size AS 文件字节, char_count AS 字数, page_count AS PDF页数, "
                    "uploaded_at AS 导入时间 FROM books ORDER BY category, title", conn)
            finally:
                conn.close()

            if books_df.empty:
                st.info("图书库暂无图书。")
            else:
                display_df = books_df.copy()
                display_df["文件大小"] = display_df["文件字节"].apply(format_file_size)
                display_df = display_df.drop(columns=["文件字节"])
                st.dataframe(display_df, hide_index=True)

                st.markdown("### 删除图书")
                del_book_id = st.number_input(
                    "输入要删除的图书 ID", min_value=1, step=1, key="book_delete_id")
                if st.button("🗑️ 删除此图书", key="book_delete_run"):
                    conn = get_connection()
                    try:
                        conn.execute(
                            "DELETE FROM reading_progress WHERE book_id=?", (del_book_id,))
                        conn.execute(
                            "DELETE FROM book_files WHERE book_id=?", (del_book_id,))
                        cur = conn.execute("DELETE FROM books WHERE id=?", (del_book_id,))
                        conn.commit()
                        deleted = cur.rowcount
                    finally:
                        conn.close()
                    clear_search_caches()
                    backup_book_library()
                    if deleted:
                        st.success("图书已删除。")
                        st.rerun()
                    else:
                        st.warning("未找到该 ID 的图书。")

            st.caption("旧版 DOC/WPS 二进制文件若无法提取文字，会保存原文件，可在阅读页下载。")
            st.markdown("### 📦 图书库备份与恢复")
            st.caption(
                "图书库会自动备份为 book_library_backup.db；主数据库被部署包覆盖后，"
                "启动时会自动从该备份恢复。若整个运行目录被清空，请先下载备份文件保存。")
            if st.button("📦 生成图书库备份文件", key="book_backup_generate"):
                backup_book_library()
                backup_count = 0
                if os.path.exists(BOOK_BACKUP_FILE):
                    try:
                        check_conn = sqlite3.connect(BOOK_BACKUP_FILE)
                        try:
                            backup_count = check_conn.execute(
                                "SELECT COUNT(*) FROM books").fetchone()[0]
                        finally:
                            check_conn.close()
                    except Exception:
                        backup_count = 0
                if backup_count:
                    with open(BOOK_BACKUP_FILE, "rb") as f:
                        st.session_state["book_backup_export"] = f.read()
                    st.success("✅ 备份已生成，可点击下方按钮下载。")
                else:
                    st.warning("当前没有可备份的图书。")
            export_data = st.session_state.get("book_backup_export")
            if export_data:
                st.download_button(
                    "⬇️ 下载图书库备份（.db）",
                    data=export_data,
                    file_name="book_library_backup.db",
                    mime="application/octet-stream",
                    key="book_backup_download")
                if st.button("清除导出缓存", key="book_backup_export_clear"):
                    st.session_state.pop("book_backup_export", None)
                    st.rerun()

            st.markdown("#### 📥 恢复图书库备份")
            st.caption("上传之前下载的 .db 备份，会覆盖当前图书库，请确认后再操作。")
            uploaded_backup = st.file_uploader(
                "选择图书库备份文件（.db）", type=["db"], key="book_backup_upload")
            if st.button("恢复此备份", key="book_backup_restore"):
                if not uploaded_backup:
                    st.warning("请先选择备份文件。")
                else:
                    raw = uploaded_backup.getvalue()
                    if not raw.startswith(b"SQLite format 3\x00"):
                        st.error("不是有效的 SQLite 备份文件。")
                    else:
                        temp_path = BOOK_BACKUP_FILE + ".upload"
                        with open(temp_path, "wb") as f:
                            f.write(raw)
                        try:
                            test_conn = sqlite3.connect(temp_path)
                            try:
                                count = test_conn.execute(
                                    "SELECT COUNT(*) FROM books").fetchone()[0]
                            finally:
                                test_conn.close()
                        except Exception:
                            count = 0
                        if not count:
                            st.error("备份中没有图书数据或文件结构不正确。")
                            try:
                                os.remove(temp_path)
                            except Exception:
                                pass
                        else:
                            try:
                                os.replace(temp_path, BOOK_BACKUP_FILE)
                            except Exception as e:
                                st.error(f"保存备份失败：{e}")
                            else:
                                if restore_book_library_from_backup(force=True):
                                    clear_search_caches()
                                    st.success("✅ 图书库已恢复。")
                                    st.rerun()
                                else:
                                    st.error("恢复失败，请检查备份文件。")

        with tab_password:
            st.subheader("🔐 后台登录密码")
            st.caption("修改后立即生效，下次进入后台请使用新密码。密码以哈希形式保存，不保存明文。")
            with st.form("change_admin_pwd_form"):
                current_pwd = st.text_input("当前密码", type="password")
                new_pwd = st.text_input("新密码（至少 4 位）", type="password")
                confirm_pwd = st.text_input("确认新密码", type="password")
                change_submitted = st.form_submit_button("💾 保存新密码", type="primary")
            if change_submitted:
                if not verify_admin_password(current_pwd):
                    st.error("当前密码不正确。")
                elif len(new_pwd) < 4:
                    st.error("新密码至少 4 位。")
                elif new_pwd != confirm_pwd:
                    st.error("两次输入的新密码不一致。")
                else:
                    ok, error = set_admin_password(new_pwd)
                    if ok:
                        st.success("✅ 后台密码已修改，下次登录请使用新密码。")
                    else:
                        st.error(f"密码修改失败：{error}")
            st.caption("如忘记密码，可在数据库中删除 system_config 表的 admin_password_salt、admin_password_hash 两行，系统会恢复默认密码 123。")

            st.divider()
            st.subheader("📥 下载原文件密码")
            st.caption("阅读页下载图书原文件时需要此密码；默认 123，可在此修改。")
            with st.form("change_download_pwd_form"):
                current_dl_pwd = st.text_input("当前下载密码", type="password")
                new_dl_pwd = st.text_input("新下载密码（至少 4 位）", type="password")
                confirm_dl_pwd = st.text_input("确认新下载密码", type="password")
                dl_change_submitted = st.form_submit_button("💾 保存下载密码", type="primary")
            if dl_change_submitted:
                if not verify_download_password(current_dl_pwd):
                    st.error("当前下载密码不正确。")
                elif len(new_dl_pwd) < 4:
                    st.error("新下载密码至少 4 位。")
                elif new_dl_pwd != confirm_dl_pwd:
                    st.error("两次输入的新下载密码不一致。")
                else:
                    ok, error = set_download_password(new_dl_pwd)
                    if ok:
                        st.success("✅ 下载密码已修改。")
                    else:
                        st.error(f"下载密码修改失败：{error}")
            st.caption("如忘记下载密码，可删除 system_config 表的 download_password_salt、download_password_hash 两行，恢复默认 123。")

        with tab_danger:
            st.subheader("☢️ 高危指令区")
            clear_table = st.selectbox(
                "选择要物理清空的数据表",
                [FORMULA_TABLE, RECIPE_TABLE, VERIF_TABLE, "query_logs"],
                key="admin_clear_table",
            )
            confirm_text = st.text_input(
                f"如果您确切明白后果，请在此处手动输入：确认清空{clear_table}", key="admin_clear_confirm")
            if st.button("🚨 强制执行一键清空", key="admin_clear_run"):
                if confirm_text == f"确认清空{clear_table}":
                    conn = get_connection()
                    conn.execute(f'DELETE FROM "{clear_table}"')
                    conn.commit()
                    conn.close()
                    clear_search_caches()
                    st.success(f"💀 核心指令已下达，【{clear_table}】已被永久清空。")
                else:
                    st.warning("防误触锁生效：口令不正确。")
